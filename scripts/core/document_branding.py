# -*- coding: utf-8 -*-
"""
ITInfra Business Ops — Subsystem Brand Identity and Document Styling (SPEC-16)
Gestione centralizzata dell'identita visiva aziendale, asset logo, palette colori
e funzioni di decorazione per documenti Word (.docx), PDF e HTML.
"""

import base64
from pathlib import Path
from typing import Dict, Any, Optional

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.platypus import Table, TableStyle, Image as RLImage, Paragraph
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

REPO_ROOT = Path(__file__).resolve().parent.parent.parent


class BrandConfig:
    """Parametri anagrafici e grafici ufficiali di Aure System."""
    COMPANY_NAME = "Aure System di Eduardo Possumato"
    LEGAL_SIGNATURE = "Aure System di Eduardo Possumato"
    ADDRESS = "Via Luigi Tansillo, 54 F - 80125 Napoli (NA)"
    VAT_ID = "IT07714231219"
    FISCAL_CODE = "IT07714231219"
    PHONE = "+39 333 7328065"
    PEC = "salvatorepossumato@pec.it"
    EMAIL = "info@auresystem.it"
    WEBSITE = "www.auresystem.it"

    # Palette Cromatica Ufficiale
    HEX_PRIMARY = "#1E3A8A"     # Aure Deep Navy
    HEX_SECONDARY = "#2563EB"   # Tech Accent Blue
    HEX_TEXT = "#0F172A"        # Slate Dark
    HEX_MUTED = "#64748B"       # Slate Gray
    HEX_SURFACE = "#F8FAFC"     # Soft Light Gray
    HEX_BORDER = "#CBD5E1"      # Gray Border
    HEX_ACCENT = "#B45309"      # Amber Warning/Option

    RGB_PRIMARY = RGBColor(0x1E, 0x3A, 0x8A)
    RGB_SECONDARY = RGBColor(0x25, 0x63, 0xEB)
    RGB_TEXT = RGBColor(0x0F, 0x17, 0x2A)
    RGB_MUTED = RGBColor(0x64, 0x74, 0x8B)
    RGB_BORDER = RGBColor(0xCB, 0xD5, 0xE1)

    @classmethod
    def get_logo_path(cls) -> Optional[Path]:
        cand = REPO_ROOT / "templates" / "assets" / "brand" / "logo.png"
        if cand.is_file():
            return cand
        alt = Path(r"C:\Users\auresystem\repos\itinfra\templates\assets\brand\logo.png")
        if alt.is_file():
            return alt
        return None

    @classmethod
    def get_logo_base64(cls) -> str:
        b64_file = REPO_ROOT / "templates" / "assets" / "brand" / "logo.base64.txt"
        if b64_file.is_file():
            return b64_file.read_text(encoding="utf-8").strip()
        logo_path = cls.get_logo_path()
        if logo_path and logo_path.is_file():
            raw = logo_path.read_bytes()
            return base64.b64encode(raw).decode("utf-8")
        return ""


class DocumentBranding:
    """Motore di applicazione del branding per documenti DOCX, PDF e HTML."""

    @staticmethod
    def apply_docx_margins(doc: docx.Document, top=0.8, bottom=0.8, left=0.8, right=0.8):
        """Imposta margini standard A4 in pollici."""
        for sec in doc.sections:
            sec.top_margin = Inches(top)
            sec.bottom_margin = Inches(bottom)
            sec.left_margin = Inches(left)
            sec.right_margin = Inches(right)

    @staticmethod
    def add_docx_branded_header(doc: docx.Document, doc_title: str, doc_subtitle: str = ""):
        """Inserisce la testata istituzionale ufficiale con logo a sinistra e dati fornitore a destra."""
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Cella Sinistra: Logo
        cell_left = table.cell(0, 0)
        p_logo = cell_left.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        logo_path = BrandConfig.get_logo_path()
        if logo_path and logo_path.is_file():
            try:
                p_logo.add_run().add_picture(str(logo_path), width=Inches(1.25))
            except Exception:
                r_fallback = p_logo.add_run(BrandConfig.COMPANY_NAME)
                r_fallback.bold = True
                r_fallback.font.color.rgb = BrandConfig.RGB_PRIMARY

        # Cella Destra: Anagrafica Aziendale Fornitore
        cell_right = table.cell(0, 1)
        p_right = cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_comp = p_right.add_run(f"{BrandConfig.COMPANY_NAME}\n")
        r_comp.bold = True
        r_comp.font.size = Pt(9.5)
        r_comp.font.color.rgb = BrandConfig.RGB_PRIMARY

        r_info = p_right.add_run(
            f"{BrandConfig.ADDRESS}\n"
            f"P.IVA: {BrandConfig.VAT_ID}  |  PEC: {BrandConfig.PEC}\n"
            f"Tel: {BrandConfig.PHONE}  |  Web: {BrandConfig.WEBSITE}"
        )
        r_info.font.size = Pt(8.5)
        r_info.font.color.rgb = BrandConfig.RGB_MUTED

        # Spaziatura e Titolo Documento
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(8)

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t = p_title.add_run(doc_title)
        r_t.bold = True
        r_t.font.size = Pt(16)
        r_t.font.color.rgb = BrandConfig.RGB_PRIMARY

        if doc_subtitle:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_s = p_sub.add_run(doc_subtitle)
            r_s.font.size = Pt(9.5)
            r_s.font.color.rgb = BrandConfig.RGB_MUTED

    @staticmethod
    def add_docx_branded_footer(doc: docx.Document):
        """Aggiunge pie di pagina con nota legale di riservatezza."""
        for sec in doc.sections:
            footer = sec.footer
            p_f = footer.paragraphs[0]
            p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_f = p_f.add_run(
                f"{BrandConfig.COMPANY_NAME} — Documento Riservato ai sensi del Reg. UE 2016/679 (GDPR)"
            )
            r_f.font.size = Pt(8)
            r_f.font.color.rgb = BrandConfig.RGB_MUTED

    @staticmethod
    def get_pdf_branded_header(doc_title: str, doc_id: str, date_str: str, valid_until: str = "") -> Table:
        """Restituisce una Tabella Flowable ReportLab pronta per la testa di pagina del PDF."""
        styles = getSampleStyleSheet()
        normal = styles['Normal']
        
        info_html = f"<font size='10' color='{BrandConfig.HEX_PRIMARY}'><b>{BrandConfig.COMPANY_NAME}</b></font><br/>" \
                    f"<font size='8' color='{BrandConfig.HEX_MUTED}'>{BrandConfig.ADDRESS}<br/>" \
                    f"P.IVA: <b>{BrandConfig.VAT_ID}</b>  |  Tel: {BrandConfig.PHONE}<br/>" \
                    f"PEC: {BrandConfig.PEC}  |  {BrandConfig.WEBSITE}</font>"
        p_info = Paragraph(info_html, normal)

        logo_path = BrandConfig.get_logo_path()
        if logo_path and logo_path.is_file():
            img = RLImage(str(logo_path), width=70, height=70)
        else:
            img = Paragraph(f"<b>{BrandConfig.COMPANY_NAME}</b>", normal)

        header_table = Table([[img, p_info]], colWidths=[85, 435])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LINEBELOW', (0, 0), (-1, -1), 1, colors.HexColor(BrandConfig.HEX_SECONDARY)),
        ]))
        return header_table

    @staticmethod
    def get_html_print_css() -> str:
        """CSS Print-Ready per formati A4 e rendering Zero-CDN responsive."""
        return f"""
        @page {{
            size: A4 portrait;
            margin: 15mm 18mm 15mm 18mm;
        }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            color: {BrandConfig.HEX_TEXT};
            background: #ffffff;
            line-height: 1.5;
            margin: 25px;
        }}
        .brand-header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            border-bottom: 2px solid {BrandConfig.HEX_SECONDARY};
            padding-bottom: 14px;
            margin-bottom: 20px;
        }}
        .brand-logo-img {{
            height: 65px;
            width: auto;
            max-width: 140px;
            object-fit: contain;
        }}
        .brand-company-info {{
            text-align: right;
            font-size: 11.5px;
            color: {BrandConfig.HEX_MUTED};
            line-height: 1.4;
        }}
        .brand-company-name {{
            font-size: 14px;
            font-weight: 700;
            color: {BrandConfig.HEX_PRIMARY};
            margin-bottom: 2px;
        }}
        .doc-title-box {{
            text-align: center;
            margin-bottom: 24px;
        }}
        .doc-title-box h1 {{
            margin: 0;
            font-size: 20px;
            color: {BrandConfig.HEX_PRIMARY};
            text-transform: uppercase;
            letter-spacing: 0.5px;
        }}
        .doc-meta {{
            font-size: 13px;
            color: {BrandConfig.HEX_MUTED};
            margin-top: 4px;
        }}
        .card-grid {{
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 20px;
            margin-bottom: 20px;
        }}
        .party-card {{
            background: {BrandConfig.HEX_SURFACE};
            border: 1px solid {BrandConfig.HEX_BORDER};
            border-radius: 6px;
            padding: 14px;
            font-size: 13px;
        }}
        .party-card-title {{
            font-size: 11.5px;
            font-weight: 700;
            text-transform: uppercase;
            color: {BrandConfig.HEX_PRIMARY};
            border-bottom: 1px solid {BrandConfig.HEX_BORDER};
            padding-bottom: 4px;
            margin-bottom: 8px;
        }}
        .data-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 12.5px;
            margin-top: 14px;
            margin-bottom: 18px;
        }}
        .data-table th {{
            background: {BrandConfig.HEX_PRIMARY};
            color: #ffffff;
            font-weight: 600;
            padding: 8px 10px;
            text-align: left;
        }}
        .data-table td {{
            padding: 8px 10px;
            border-bottom: 1px solid #e2e8f0;
        }}
        .data-table tr:nth-child(even) {{
            background: {BrandConfig.HEX_SURFACE};
        }}
        .text-right {{ text-align: right; }}
        .text-center {{ text-align: center; }}
        .totals-card {{
            width: 320px;
            margin-left: auto;
            background: {BrandConfig.HEX_SURFACE};
            border: 1px solid {BrandConfig.HEX_BORDER};
            border-radius: 6px;
            padding: 14px;
            margin-top: 20px;
            font-size: 13px;
        }}
        .grand-total {{
            font-size: 16px;
            font-weight: 700;
            color: {BrandConfig.HEX_PRIMARY};
            border-top: 2px solid {BrandConfig.HEX_SECONDARY};
            padding-top: 8px;
            margin-top: 8px;
            display: flex;
            justify-content: space-between;
        }}
        .sign-grid {{
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 40px;
            margin-top: 40px;
        }}
        .sign-box {{
            border-top: 1px dashed {BrandConfig.HEX_MUTED};
            padding-top: 8px;
            text-align: center;
            font-size: 12px;
            color: {BrandConfig.HEX_MUTED};
            margin-top: 45px;
        }}
        .footer-legal {{
            text-align: center;
            font-size: 10px;
            color: {BrandConfig.HEX_MUTED};
            border-top: 1px solid #e2e8f0;
            padding-top: 12px;
            margin-top: 35px;
        }}
        @media print {{
            body {{ margin: 0; }}
            .no-print {{ display: none; }}
        }}
        """
