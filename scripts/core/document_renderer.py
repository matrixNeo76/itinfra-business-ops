import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Dict, List, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

class DocumentRenderer:
    """
    Motore universale di visualizzazione ed esportazione documentale:
    - Conversione da Fattura Elettronica XML (SDI v1.2 FPR12) ad HTML e PDF di cortesia
    - Esportazione di Preventivi ed Offerte in formato Microsoft Word (.docx) e PDF
    """

    @staticmethod
    def render_xml_invoice_to_html(xml_path: Path, out_html_path: Optional[Path] = None) -> Path:
        """Legge un file XML FatturaPA SDI FPR12 e genera una vista grafica HTML human-readable."""
        tree = ET.parse(xml_path)
        root = tree.getroot()

        # Namespace cleaner
        def find_text(elem, *tags):
            curr = elem
            for t in tags:
                if curr is None:
                    return ""
                found = None
                for child in curr:
                    tag_name = child.tag.split("}")[-1]
                    if tag_name == t:
                        found = child
                        break
                curr = found
            return curr.text.strip() if curr is not None and curr.text else ""

        # Dati Trasmissione
        dati_trasm = None
        cedente = None
        cessionario = None
        dati_doc = None
        dati_beni = None
        dati_pag = None

        for child in root:
            tag = child.tag.split("}")[-1]
            if tag == "FatturaElettronicaHeader":
                for h_child in child:
                    htag = h_child.tag.split("}")[-1]
                    if htag == "DatiTrasmissione":
                        dati_trasm = h_child
                    elif htag == "CedentePrestatore":
                        cedente = h_child
                    elif htag == "CessionarioCommittente":
                        cessionario = h_child
            elif tag == "FatturaElettronicaBody":
                for b_child in child:
                    btag = b_child.tag.split("}")[-1]
                    if btag == "DatiGenerali":
                        for dg_child in b_child:
                            if dg_child.tag.split("}")[-1] == "DatiGeneraliDocumento":
                                dati_doc = dg_child
                    elif btag == "DatiBeniServizi":
                        dati_beni = b_child
                    elif btag == "DatiPagamento":
                        dati_pag = b_child

        # Estrazione Campi Cedente
        ced_nome = find_text(cedente, "DatiAnagrafici", "Anagrafica", "Denominazione")
        if not ced_nome:
            ced_nome = f"{find_text(cedente, 'DatiAnagrafici', 'Anagrafica', 'Nome')} {find_text(cedente, 'DatiAnagrafici', 'Anagrafica', 'Cognome')}".strip()
        ced_piva = find_text(cedente, "DatiAnagrafici", "IdFiscaleIVA", "IdCodice")
        ced_cf = find_text(cedente, "DatiAnagrafici", "CodiceFiscale")
        ced_via = find_text(cedente, "Sede", "Indirizzo")
        ced_cap = find_text(cedente, "Sede", "CAP")
        ced_comune = find_text(cedente, "Sede", "Comune")
        ced_prov = find_text(cedente, "Sede", "Provincia")

        # Estrazione Campi Cessionario
        ces_nome = find_text(cessionario, "DatiAnagrafici", "Anagrafica", "Denominazione")
        if not ces_nome:
            ces_nome = f"{find_text(cessionario, 'DatiAnagrafici', 'Anagrafica', 'Nome')} {find_text(cessionario, 'DatiAnagrafici', 'Anagrafica', 'Cognome')}".strip()
        ces_piva = find_text(cessionario, "DatiAnagrafici", "IdFiscaleIVA", "IdCodice")
        ces_cf = find_text(cessionario, "DatiAnagrafici", "CodiceFiscale")
        ces_via = find_text(cessionario, "Sede", "Indirizzo")
        ces_cap = find_text(cessionario, "Sede", "CAP")
        ces_comune = find_text(cessionario, "Sede", "Comune")
        ces_prov = find_text(cessionario, "Sede", "Provincia")

        # Dati Documento
        doc_tipo = find_text(dati_doc, "TipoDocumento") or "TD01"
        doc_divisa = find_text(dati_doc, "Divisa") or "EUR"
        doc_data = find_text(dati_doc, "Data")
        doc_numero = find_text(dati_doc, "Numero")
        doc_totale = float(find_text(dati_doc, "ImportoTotaleDocumento") or 0.0)

        cod_sdi = find_text(dati_trasm, "CodiceDestinatario")
        pec_dest = find_text(dati_trasm, "PECDestinatario")

        # Righe di dettaglio
        lines_html = []
        if dati_beni:
            for item in dati_beni:
                if item.tag.split("}")[-1] == "DettaglioLinee":
                    num_lin = find_text(item, "NumeroLinea")
                    desc = find_text(item, "Descrizione")
                    qty = find_text(item, "Quantita") or "1"
                    prz_u = float(find_text(item, "PrezzoUnitario") or 0.0)
                    prz_t = float(find_text(item, "PrezzoTotale") or 0.0)
                    iva = find_text(item, "AliquotaIVA") or "22.00"
                    lines_html.append(f"""
                    <tr>
                      <td style="padding:10px; border-bottom:1px solid #e2e8f0; text-align:center;">{num_lin}</td>
                      <td style="padding:10px; border-bottom:1px solid #e2e8f0;">{desc}</td>
                      <td style="padding:10px; border-bottom:1px solid #e2e8f0; text-align:center;">{qty}</td>
                      <td style="padding:10px; border-bottom:1px solid #e2e8f0; text-align:right;">€ {prz_u:.2f}</td>
                      <td style="padding:10px; border-bottom:1px solid #e2e8f0; text-align:center;">{float(iva):.0f}%</td>
                      <td style="padding:10px; border-bottom:1px solid #e2e8f0; text-align:right; font-weight:600;">€ {prz_t:.2f}</td>
                    </tr>
                    """)

        # Riepilogo IVA
        riepilogo_html = []
        if dati_beni:
            for item in dati_beni:
                if item.tag.split("}")[-1] == "DatiRiepilogo":
                    al_iva = find_text(item, "AliquotaIVA")
                    imp_imp = float(find_text(item, "ImponibileImporto") or 0.0)
                    imposta = float(find_text(item, "Imposta") or 0.0)
                    riepilogo_html.append(f"""
                    <div style="display:flex; justify-content:space-between; margin-bottom:6px; font-size:14px;">
                      <span>Imponibile IVA {float(al_iva):.0f}%:</span>
                      <strong>€ {imp_imp:.2f}</strong>
                    </div>
                    <div style="display:flex; justify-content:space-between; margin-bottom:6px; font-size:14px;">
                      <span>Imposta IVA {float(al_iva):.0f}%:</span>
                      <strong>€ {imposta:.2f}</strong>
                    </div>
                    """)

        # Dati Pagamento
        pag_info_html = ""
        if dati_pag:
            for dp in dati_pag:
                if dp.tag.split("}")[-1] == "DettaglioPagamento":
                    mod = find_text(dp, "ModalitaPagamento") or "MP05 (Bonifico)"
                    dt_scad = find_text(dp, "DataScadenzaPagamento")
                    imp_scad = float(find_text(dp, "ImportoPagamento") or 0.0)
                    iban_val = find_text(dp, "IBAN")
                    pag_info_html += f"""
                    <div style="margin-top:6px;">
                      • Rata Scadenza: <strong>{dt_scad}</strong> — Importo: <strong>€ {imp_scad:.2f}</strong> (Mod: {mod})
                    </div>
                    """
            if iban_val:
                pag_info_html += f"<div style='margin-top:8px;'>Coordinate IBAN: <code>{iban_val}</code></div>"

        html = f"""<!DOCTYPE html>
<html lang="it">
<head>
<meta charset="UTF-8">
<title>Fattura Elettronica {doc_numero} — {ces_nome}</title>
<style>
  body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif; margin: 30px; color: #1e293b; background: #fff; }}
  .invoice-card {{ max-width: 900px; margin: 0 auto; border: 1px solid #cbd5e1; border-radius: 8px; padding: 32px; box-shadow: 0 4px 6px -1px rgb(0 0 0 / 0.1); }}
  .header {{ display: flex; justify-content: space-between; border-bottom: 2px solid #2563eb; padding-bottom: 20px; margin-bottom: 24px; }}
  .box-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 24px; margin-bottom: 24px; }}
  .party-box {{ border: 1px solid #e2e8f0; border-radius: 6px; padding: 16px; background: #f8fafc; font-size: 13px; line-height: 1.6; }}
  .party-title {{ font-size: 14px; font-weight: bold; color: #1e40af; text-transform: uppercase; margin-bottom: 8px; border-bottom: 1px solid #e2e8f0; padding-bottom: 4px; }}
  .table {{ width: 100%; border-collapse: collapse; margin-top: 16px; font-size: 13px; }}
  .table th {{ background: #1e40af; color: white; padding: 10px; text-align: left; }}
  .summary-box {{ display: flex; justify-content: flex-end; margin-top: 24px; }}
  .summary-inner {{ width: 340px; border: 1px solid #cbd5e1; border-radius: 6px; padding: 16px; background: #f8fafc; }}
  .grand-total {{ font-size: 18px; font-weight: bold; color: #1e3a8a; border-top: 2px solid #2563eb; padding-top: 8px; margin-top: 8px; display: flex; justify-content: space-between; }}
  .footer-notice {{ margin-top: 30px; padding: 12px; background: #eff6ff; border-left: 4px solid #2563eb; font-size: 12px; color: #1e40af; border-radius: 0 4px 4px 0; }}
  @media print {{ body {{ margin: 0; }} .invoice-card {{ border: none; box-shadow: none; padding: 0; }} }}
</style>
</head>
<body>
<div class="invoice-card">
  <div class="header">
    <div>
      <h1 style="margin:0; font-size:24px; color:#1e3a8a;">FATTURA ELETTRONICA</h1>
      <div style="font-size:14px; color:#64748b; margin-top:4px;">Tipo Documento: <strong>{doc_tipo}</strong> | Formato: <strong>FPR12 (B2B)</strong></div>
      <div style="font-size:15px; font-weight:bold; margin-top:6px; color:#0f172a;">Numero: {doc_numero}</div>
      <div style="font-size:13px; color:#64748b;">Data Emissione: {doc_data}</div>
    </div>
    <div style="text-align:right; font-size:13px;">
      <span style="display:inline-block; padding:4px 10px; background:#dbeafe; color:#1e40af; border-radius:4px; font-weight:600;">COPIA DI CORTESIA</span>
      <div style="margin-top:10px; color:#64748b;">Codice SDI: <strong>{cod_sdi}</strong></div>
      {f"<div style='color:#64748b;'>PEC: <strong>{pec_dest}</strong></div>" if pec_dest else ""}
    </div>
  </div>

  <div class="box-grid">
    <div class="party-box">
      <div class="party-title">Cedente / Prestatore (Fornitore)</div>
      <strong style="font-size:14px; color:#0f172a;">{ced_nome}</strong><br>
      {ced_via}<br>
      {ced_cap} {ced_comune} ({ced_prov})<br>
      P.IVA: <strong>{ced_piva}</strong><br>
      {f"Codice Fiscale: {ced_cf}<br>" if ced_cf and ced_cf != ced_piva else ""}
    </div>
    <div class="party-box">
      <div class="party-title">Cessionario / Committente (Cliente)</div>
      <strong style="font-size:14px; color:#0f172a;">{ces_nome}</strong><br>
      {ces_via}<br>
      {ces_cap} {ces_comune} ({ces_prov})<br>
      P.IVA: <strong>{ces_piva}</strong><br>
      {f"Codice Fiscale: {ces_cf}<br>" if ces_cf else ""}
    </div>
  </div>

  <table class="table">
    <thead>
      <tr>
        <th style="width:50px; text-align:center;">N°</th>
        <th>Descrizione Beni / Servizi</th>
        <th style="width:60px; text-align:center;">Q.tà</th>
        <th style="width:110px; text-align:right;">Prezzo Unit.</th>
        <th style="width:70px; text-align:center;">IVA</th>
        <th style="width:120px; text-align:right;">Importo</th>
      </tr>
    </thead>
    <tbody>
      {"".join(lines_html)}
    </tbody>
  </table>

  <div class="summary-box">
    <div class="summary-inner">
      {"".join(riepilogo_html)}
      <div class="grand-total">
        <span>TOTALE DOCUMENTO:</span>
        <span>€ {doc_totale:.2f}</span>
      </div>
    </div>
  </div>

  <div class="footer-notice">
    <strong>Condizioni di Pagamento & Scadenze:</strong>
    {pag_info_html}
    <div style="margin-top:10px; font-size:11px; color:#64748b; border-top:1px solid #cbd5e1; padding-top:6px;">
      Documento privo di valenza fiscale ai sensi dell'art. 21 D.P.R. 633/72. L'originale della fattura elettronica è disponibile nella Sua area riservata dell'Agenzia delle Entrate o recapitato via SDI al codice {cod_sdi}.
    </div>
  </div>
</div>
</body>
</html>
"""
        target_path = out_html_path or xml_path.with_suffix(".html")
        target_path.write_text(html, encoding="utf-8")
        return target_path

    @staticmethod
    def render_quote_to_docx(quote_data: Dict[str, Any], out_docx_path: Path) -> Path:
        """Esporta il preventivo in formato professionale Microsoft Word (.docx)."""
        doc = docx.Document()

        # Impostazione margini 2 cm
        for sec in doc.sections:
            sec.top_margin = Inches(0.8)
            sec.bottom_margin = Inches(0.8)
            sec.left_margin = Inches(0.8)
            sec.right_margin = Inches(0.8)

        # Stili e Titolo
        title = doc.add_heading("PROPOSTA COMMERCIALE & PREVENTIVO", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_sub = p_sub.add_run(f"ID Documento: {quote_data.get('quote_id')}  |  Data: {quote_data.get('created_at')}  |  Validità: fino al {quote_data.get('valid_until')}")
        r_sub.font.size = Pt(10)
        r_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        doc.add_paragraph()

        # Tabella Dati Intestazione Cliente / Fornitore
        head_table = doc.add_table(rows=1, cols=2)
        head_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        head_table.autofit = True

        cell_f = head_table.cell(0, 0)
        p_f = cell_f.paragraphs[0]
        p_f.add_run("SOCIETÀ FORNITRICE:\n").bold = True
        p_f.add_run("Aure System / ITInfra Business Ops\nVia Luigi Tansillo, 54 F\n80125 Napoli (NA)\nP.IVA: IT07714231219")

        cell_c = head_table.cell(0, 1)
        p_c = cell_c.paragraphs[0]
        p_c.add_run("SPETTABILE COMMITTENTE:\n").bold = True
        p_c.add_run(f"{quote_data.get('client_name', 'Cliente')}\n")
        if quote_data.get('client_contact'):
            p_c.add_run(f"C.a.: {quote_data['client_contact']}\n")
        p_c.add_run(f"Termini di Pagamento: {quote_data.get('payment_terms', '30_60_DF_FM')}")

        doc.add_paragraph()

        # Sezioni e Tabelle per Categoria
        for cat in quote_data.get("categories", []):
            cat_name = cat.get("name", "").replace("_", " ").title()
            h_cat = doc.add_heading(cat_name, level=2)
            h_cat.runs[0].font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

            items = cat.get("items", [])
            table = doc.add_table(rows=1, cols=5)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            # Intestazioni Tabella
            hdr_cells = table.rows[0].cells
            hdr_titles = ["Codice", "Descrizione", "Q.tà", "Prezzo Unit.", "Totale Netto"]
            for i, text in enumerate(hdr_titles):
                hdr_cells[i].text = text
                hdr_cells[i].paragraphs[0].runs[0].bold = True
                hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)

            for it in items:
                row_cells = table.add_row().cells
                row_cells[0].text = it.get("part_number", "")
                
                desc_text = it.get("description", "")
                if it.get("is_optional"):
                    desc_text += "  [OPZIONE]"
                row_cells[1].text = desc_text

                row_cells[2].text = str(it.get("quantity", 1))
                row_cells[3].text = f"€ {it.get('unit_price', 0.0):.2f}"
                row_cells[4].text = f"€ {it.get('line_total', 0.0):.2f}"

                for cell in row_cells:
                    cell.paragraphs[0].runs[0].font.size = Pt(9)

            doc.add_paragraph()

        # Box Totali
        tot_net = quote_data.get("totals", {}).get("total_net", 0.0)
        tot_vat = round(tot_net * 0.22, 2)
        tot_gross = round(tot_net + tot_vat, 2)

        tot_table = doc.add_table(rows=3, cols=2)
        tot_table.alignment = WD_TABLE_ALIGNMENT.RIGHT
        tot_table.cell(0, 0).text = "Totale Fornitura Netta:"
        tot_table.cell(0, 1).text = f"€ {tot_net:.2f}"
        tot_table.cell(1, 0).text = "IVA (22%):"
        tot_table.cell(1, 1).text = f"€ {tot_vat:.2f}"
        tot_table.cell(2, 0).text = "TOTALE PROPOSTA:"
        tot_table.cell(2, 1).text = f"€ {tot_gross:.2f}"

        for r in tot_table.rows:
            r.cells[0].paragraphs[0].runs[0].bold = True
            r.cells[1].paragraphs[0].runs[0].bold = True

        doc.add_paragraph()

        # Condizioni
        doc.add_heading("Condizioni Generali di Fornitura", level=3)
        doc.add_paragraph(f"• Tempi di Consegna Stimati: {quote_data.get('delivery_time_weeks', 3)} settimane da conferma d'ordine.")
        doc.add_paragraph(f"• Termini di Pagamento: {quote_data.get('payment_terms', '30_60_DF_FM')}.")
        doc.add_paragraph("• Garanzia Hardware: On-site con supporto tecnico certificato come da accordi SLA.")
        doc.add_paragraph("• Le voci contrassegnate come [OPZIONE] sono attivabili su richiesta ed escluse dal totale.")

        doc.add_paragraph()

        # Firme
        sign_table = doc.add_table(rows=2, cols=2)
        sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sign_table.cell(0, 0).text = "Per la Società Fornitrice:"
        sign_table.cell(0, 1).text = "Per Accettazione Cliente (Timbro e Firma):"
        sign_table.cell(1, 0).text = "\n\n__________________________________"
        sign_table.cell(1, 1).text = f"\n\n__________________________________\n{quote_data.get('client_name')}"

        doc.save(out_docx_path)
        return out_docx_path

    @staticmethod
    def render_quote_to_pdf(quote_data: Dict[str, Any], out_pdf_path: Path) -> Path:
        """Esporta il preventivo in formato PDF stampabile ad alta risoluzione."""
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        doc = SimpleDocTemplate(
            str(out_pdf_path),
            pagesize=A4,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1E3A8A'),
            alignment=1
        )
        subtitle_style = ParagraphStyle(
            'DocSub',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#64748B'),
            alignment=1
        )
        h2_style = ParagraphStyle(
            'CategoryHeader',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=12,
            leading=16,
            textColor=colors.HexColor('#2563EB'),
            spaceBefore=14,
            spaceAfter=6
        )
        normal_style = ParagraphStyle(
            'NormalText',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=11
        )
        bold_style = ParagraphStyle(
            'BoldText',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8.5,
            leading=11
        )

        story = []

        # Titolo e Sottotitolo
        story.append(Paragraph("PROPOSTA COMMERCIALE & PREVENTIVO", title_style))
        story.append(Paragraph(f"ID Proposta: <b>{quote_data.get('quote_id')}</b>  |  Data: {quote_data.get('created_at')}  |  Validità: fino al {quote_data.get('valid_until')}", subtitle_style))
        story.append(Spacer(1, 14))

        # Intestazione Fornitore / Cliente
        client_contact = f"<br/>C.a.: <b>{quote_data['client_contact']}</b>" if quote_data.get('client_contact') else ""
        parties_data = [
            [
                Paragraph("<b>FORNITORE:</b><br/>Aure System / ITInfra Business Ops<br/>Via Luigi Tansillo, 54 F - 80125 Napoli (NA)<br/>P.IVA: IT07714231219", normal_style),
                Paragraph(f"<b>COMMITTENTE:</b><br/><b>{quote_data.get('client_name', 'Cliente')}</b>{client_contact}<br/>Termini Pagamento: {quote_data.get('payment_terms', '30_60_DF_FM')}", normal_style)
            ]
        ]
        t_parties = Table(parties_data, colWidths=[260, 260])
        t_parties.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
            ('RIGHTPADDING', (0, 0), (-1, -1), 10),
        ]))
        story.append(t_parties)
        story.append(Spacer(1, 10))

        # Tabelle Categorie
        for cat in quote_data.get("categories", []):
            cat_name = cat.get("name", "").replace("_", " ").title()
            story.append(Paragraph(cat_name, h2_style))

            table_data = [[
                Paragraph("<b>Codice</b>", bold_style),
                Paragraph("<b>Descrizione Articolo / Servizio</b>", bold_style),
                Paragraph("<b>Q.tà</b>", bold_style),
                Paragraph("<b>Prezzo Unit.</b>", bold_style),
                Paragraph("<b>Totale Netto</b>", bold_style)
            ]]

            for it in cat.get("items", []):
                desc = it.get("description", "")
                if it.get("is_optional"):
                    desc += " <b><font color='#b45309'>[OPZIONE]</font></b>"
                table_data.append([
                    Paragraph(f"<code>{it.get('part_number', '')}</code>", normal_style),
                    Paragraph(desc, normal_style),
                    Paragraph(str(it.get("quantity", 1)), normal_style),
                    Paragraph(f"€ {it.get('unit_price', 0.0):.2f}", normal_style),
                    Paragraph(f"€ {it.get('line_total', 0.0):.2f}", bold_style)
                ])

            t_cat = Table(table_data, colWidths=[80, 250, 40, 75, 75])
            t_cat.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
                ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ('LEFTPADDING', (0, 0), (-1, -1), 5),
                ('RIGHTPADDING', (0, 0), (-1, -1), 5),
                ('ALIGN', (2, 0), (2, -1), 'CENTER'),
                ('ALIGN', (3, 0), (-1, -1), 'RIGHT'),
            ]))
            story.append(t_cat)
            story.append(Spacer(1, 8))

        # Totali Box
        tot_net = quote_data.get("totals", {}).get("total_net", 0.0)
        tot_vat = round(tot_net * 0.22, 2)
        tot_gross = round(tot_net + tot_vat, 2)

        tot_data = [
            [Paragraph("Totale Fornitura Netta:", normal_style), Paragraph(f"<b>€ {tot_net:.2f}</b>", normal_style)],
            [Paragraph("IVA (22%):", normal_style), Paragraph(f"<b>€ {tot_vat:.2f}</b>", normal_style)],
            [Paragraph("<b>TOTALE OFFERTA:</b>", bold_style), Paragraph(f"<b><font size='10' color='#1e3a8a'>€ {tot_gross:.2f}</font></b>", bold_style)]
        ]
        t_tot = Table(tot_data, colWidths=[140, 90], hAlign='RIGHT')
        t_tot.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ]))
        story.append(Spacer(1, 10))
        story.append(t_tot)

        # Condizioni e Firma
        story.append(Spacer(1, 12))
        terms_text = f"""<b>Condizioni di Fornitura:</b><br/>
• Tempi di Consegna: {quote_data.get('delivery_time_weeks', 3)} settimane da conferma d'ordine.<br/>
• Termini di Pagamento: {quote_data.get('payment_terms', '30_60_DF_FM')}.<br/>
• Garanzia Hardware: On-site con supporto tecnico certificato SLA.<br/>
• Le voci con tag [OPZIONE] sono escluse dall'imponibile vincolante."""
        story.append(Paragraph(terms_text, normal_style))

        sign_data = [
            [
                Paragraph("<b>Per la Società Fornitrice:</b><br/><br/><br/>__________________________________", normal_style),
                Paragraph(f"<b>Per Accettazione Cliente (Timbro e Firma):</b><br/><br/><br/>__________________________________<br/>{quote_data.get('client_name')}", normal_style)
            ]
        ]
        t_sign = Table(sign_data, colWidths=[260, 260])
        t_sign.setStyle(TableStyle([
            ('TOPPADDING', (0, 0), (-1, -1), 12),
        ]))
        story.append(Spacer(1, 20))
        story.append(t_sign)

        doc.build(story)
        return out_pdf_path

    @staticmethod
    def render_xml_invoice_to_pdf(xml_path: Path, out_pdf_path: Optional[Path] = None) -> Path:
        """Genera copia di cortesia PDF a partire dal file XML FatturaPA SDI."""
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        target_pdf = out_pdf_path or xml_path.with_suffix(".pdf")
        tree = ET.parse(xml_path)
        root = tree.getroot()

        def find_text(elem, *tags):
            curr = elem
            for t in tags:
                if curr is None:
                    return ""
                found = None
                for child in curr:
                    tag_name = child.tag.split("}")[-1]
                    if tag_name == t:
                        found = child
                        break
                curr = found
            return curr.text.strip() if curr is not None and curr.text else ""

        dati_trasm = None
        cedente = None
        cessionario = None
        dati_doc = None
        dati_beni = None
        dati_pag = None

        for child in root:
            tag = child.tag.split("}")[-1]
            if tag == "FatturaElettronicaHeader":
                for h in child:
                    htag = h.tag.split("}")[-1]
                    if htag == "DatiTrasmissione":
                        dati_trasm = h
                    elif htag == "CedentePrestatore":
                        cedente = h
                    elif htag == "CessionarioCommittente":
                        cessionario = h
            elif tag == "FatturaElettronicaBody":
                for b in child:
                    btag = b.tag.split("}")[-1]
                    if btag == "DatiGenerali":
                        for dg in b:
                            if dg.tag.split("}")[-1] == "DatiGeneraliDocumento":
                                dati_doc = dg
                    elif btag == "DatiBeniServizi":
                        dati_beni = b
                    elif btag == "DatiPagamento":
                        dati_pag = b

        ced_nome = find_text(cedente, "DatiAnagrafici", "Anagrafica", "Denominazione") or "Cedente"
        ced_piva = find_text(cedente, "DatiAnagrafici", "IdFiscaleIVA", "IdCodice")
        ced_via = find_text(cedente, "Sede", "Indirizzo")
        ced_cap = find_text(cedente, "Sede", "CAP")
        ced_comune = find_text(cedente, "Sede", "Comune")
        ced_prov = find_text(cedente, "Sede", "Provincia")

        ces_nome = find_text(cessionario, "DatiAnagrafici", "Anagrafica", "Denominazione") or "Cessionario"
        ces_piva = find_text(cessionario, "DatiAnagrafici", "IdFiscaleIVA", "IdCodice")
        ces_via = find_text(cessionario, "Sede", "Indirizzo")
        ces_cap = find_text(cessionario, "Sede", "CAP")
        ces_comune = find_text(cessionario, "Sede", "Comune")
        ces_prov = find_text(cessionario, "Sede", "Provincia")

        doc_numero = find_text(dati_doc, "Numero")
        doc_data = find_text(dati_doc, "Data")
        doc_totale = float(find_text(dati_doc, "ImportoTotaleDocumento") or 0.0)
        cod_sdi = find_text(dati_trasm, "CodiceDestinatario")

        doc_pdf = SimpleDocTemplate(
            str(target_pdf),
            pagesize=A4,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        h1_style = ParagraphStyle('InvTitle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=colors.HexColor('#1E3A8A'))
        normal_style = ParagraphStyle('InvNormal', parent=styles['Normal'], fontName='Helvetica', fontSize=8.5, leading=11)
        bold_style = ParagraphStyle('InvBold', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8.5, leading=11)

        story = []
        head_data = [
            [
                Paragraph(f"<b>FATTURA ELETTRONICA</b><br/>Numero: <b>{doc_numero}</b><br/>Data: {doc_data}", h1_style),
                Paragraph(f"<font color='#1E40AF'><b>COPIA DI CORTESIA (SDI FPR12)</b></font><br/>Codice Destinatario: <b>{cod_sdi}</b>", normal_style)
            ]
        ]
        t_head = Table(head_data, colWidths=[300, 220])
        story.append(t_head)
        story.append(Spacer(1, 14))

        parties = [
            [
                Paragraph(f"<b>CEDENTE / PRESTATORE:</b><br/><b>{ced_nome}</b><br/>{ced_via}<br/>{ced_cap} {ced_comune} ({ced_prov})<br/>P.IVA: {ced_piva}", normal_style),
                Paragraph(f"<b>CESSIONARIO / COMMITTENTE:</b><br/><b>{ces_nome}</b><br/>{ces_via}<br/>{ces_cap} {ces_comune} ({ces_prov})<br/>P.IVA: {ces_piva}", normal_style)
            ]
        ]
        t_parties = Table(parties, colWidths=[260, 260])
        t_parties.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
            ('RIGHTPADDING', (0, 0), (-1, -1), 10),
        ]))
        story.append(t_parties)
        story.append(Spacer(1, 14))

        lines_table = [[
            Paragraph("<b>N°</b>", bold_style),
            Paragraph("<b>Descrizione Beni / Servizi</b>", bold_style),
            Paragraph("<b>Q.tà</b>", bold_style),
            Paragraph("<b>Prezzo Unit.</b>", bold_style),
            Paragraph("<b>IVA</b>", bold_style),
            Paragraph("<b>Totale</b>", bold_style)
        ]]

        if dati_beni:
            for item in dati_beni:
                if item.tag.split("}")[-1] == "DettaglioLinee":
                    nl = find_text(item, "NumeroLinea")
                    ds = find_text(item, "Descrizione")
                    qt = find_text(item, "Quantita") or "1"
                    pu = float(find_text(item, "PrezzoUnitario") or 0.0)
                    pt = float(find_text(item, "PrezzoTotale") or 0.0)
                    iv = float(find_text(item, "AliquotaIVA") or 22.0)
                    lines_table.append([
                        Paragraph(nl, normal_style),
                        Paragraph(ds, normal_style),
                        Paragraph(qt, normal_style),
                        Paragraph(f"€ {pu:.2f}", normal_style),
                        Paragraph(f"{iv:.0f}%", normal_style),
                        Paragraph(f"€ {pt:.2f}", bold_style)
                    ])

        t_lines = Table(lines_table, colWidths=[30, 270, 35, 65, 45, 75])
        t_lines.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E40AF')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('ALIGN', (2, 0), (2, -1), 'CENTER'),
            ('ALIGN', (4, 0), (4, -1), 'CENTER'),
            ('ALIGN', (3, 0), (3, -1), 'RIGHT'),
            ('ALIGN', (5, 0), (5, -1), 'RIGHT'),
        ]))
        story.append(t_lines)
        story.append(Spacer(1, 14))

        tot_box = [
            [Paragraph("<b>TOTALE DOCUMENTO:</b>", bold_style), Paragraph(f"<b><font size='11' color='#1E3A8A'>€ {doc_totale:.2f}</font></b>", bold_style)]
        ]
        t_tot = Table(tot_box, colWidths=[160, 100], hAlign='RIGHT')
        t_tot.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#2563EB')),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ]))
        story.append(t_tot)

        doc_pdf.build(story)
        return target_pdf
